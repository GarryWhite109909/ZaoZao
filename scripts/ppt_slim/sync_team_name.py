#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
队员署名同步：把 pptx 里的「白明耀、易卓玥」统一改为「白明耀、易卓玥、谭依晴」。

用途：2026-09-20 谭依晴加入队伍后，把**已有成稿**（28 页原版、备份）也同步，
避免三份文件署名不一致，提交时拿错版本。

用法：
  python sync_team_name.py <file1.pptx> [file2.pptx ...]
"""

import os
import sys

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptutil import iter_shapes  # noqa: E402

OLD = "白明耀、易卓玥"
NEW = "白明耀、易卓玥、谭依晴"


def sync(path):
    prs = Presentation(path)
    hits = 0
    for i, slide in enumerate(prs.slides, 1):
        for sh in iter_shapes(slide.shapes):
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if OLD in run.text and NEW not in run.text:
                        run.text = run.text.replace(OLD, NEW)
                        hits += 1
                        print(f"  P{i:2d} {sh.name!r}: {OLD} -> {NEW}")
    if hits:
        prs.save(path)
    print(f"[{'updated' if hits else 'nochange'}] {os.path.basename(path)}  替换 {hits} 处")
    return hits


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    total = 0
    for p in sys.argv[1:]:
        if not os.path.isfile(p):
            print(f"[skip] 不存在: {p}")
            continue
        total += sync(p)
    print(f"[done] 共替换 {total} 处")


if __name__ == "__main__":
    main()
